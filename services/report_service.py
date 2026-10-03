"""Risk reporting: CSV extracts, a multi-sheet Excel workbook and a printable
PDF executive summary.

The service consumes already-computed results (it performs no risk
calculations of its own) and is independent of Streamlit, so it can be used
from scripts or scheduled jobs as well as the dashboard.
"""

from __future__ import annotations

import io
import logging
from xml.sax.saxutils import escape
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd

from analytics.backtesting import BacktestResult, BaselResult
from analytics.pnl_attribution import COMPONENT_LABELS, AttributionResult, summarise_components
from analytics.risk_limits import ConcentrationReport
from config.settings import AppSettings
from models.risk import LimitCheck
from services.market_data_service import ASSET_LABELS
from services.risk_service import METHOD_LABELS, METHODS, RiskContext, VarReport
from stress.scenario_engine import StressResult

logger = logging.getLogger(__name__)

SECTION_ORDER = (
    "Portfolio Summary",
    "Risk Metrics",
    "VaR Comparison",
    "Expected Shortfall",
    "Risk Contributions",
    "Sensitivities",
    "Stress Tests",
    "P&L Attribution",
    "VaR Backtesting",
    "Risk Limit Breaches",
)


@dataclass
class RiskReportData:
    portfolio_id: str
    portfolio_name: str
    valuation_date: date
    base_currency: str
    generated_at: datetime
    summary: pd.DataFrame
    positions: pd.DataFrame
    risk_metrics: pd.DataFrame
    var_comparison: pd.DataFrame
    expected_shortfall: pd.DataFrame
    contributions: pd.DataFrame
    contributions_by_class: pd.DataFrame
    greeks: pd.DataFrame
    bond_sensitivities: pd.DataFrame
    krd: pd.DataFrame
    stress: pd.DataFrame
    pnl_attribution: pd.DataFrame
    pnl_history: pd.DataFrame
    backtest: pd.DataFrame
    backtest_summary: pd.DataFrame
    limits: pd.DataFrame
    warnings: list[str] = field(default_factory=list)

    def tables(self) -> dict[str, pd.DataFrame]:
        """Report tables keyed by sheet / file name, in report order."""
        return {
            "Portfolio Summary": self.summary,
            "Positions": self.positions,
            "Risk Metrics": self.risk_metrics,
            "VaR Comparison": self.var_comparison,
            "Expected Shortfall": self.expected_shortfall,
            "Risk Contributions": self.contributions,
            "Contributions by Class": self.contributions_by_class,
            "Greeks": self.greeks,
            "Bond Sensitivities": self.bond_sensitivities,
            "KRD": self.krd,
            "Stress Tests": self.stress,
            "P&L Attribution": self.pnl_attribution,
            "P&L History": self.pnl_history,
            "VaR Backtest": self.backtest,
            "Backtest Summary": self.backtest_summary,
            "Risk Limits": self.limits,
        }


class ReportService:
    def __init__(self, settings: AppSettings):
        self.settings = settings

    # ------------------------------------------------------------------ build
    def build(
        self,
        ctx: RiskContext,
        var_report: VarReport,
        contributions: pd.DataFrame,
        contributions_by_class: pd.DataFrame,
        sensitivities: dict,
        stress_results: list[StressResult],
        attribution: AttributionResult,
        attribution_history: pd.DataFrame,
        backtest: BacktestResult,
        limit_checks: list[LimitCheck],
        concentration: ConcentrationReport,
        portfolio_volatility: float,
        basel: BaselResult | None = None,
    ) -> RiskReportData:
        nav = ctx.nav
        cfg = var_report.config
        hist99 = var_report.get("HISTORICAL", 0.99)
        worst = min(stress_results, key=lambda r: r.total_pnl) if stress_results else None
        summary = pd.DataFrame(
            [
                ("Portfolio", f"{ctx.portfolio.name} ({ctx.portfolio.portfolio_id})"),
                ("Valuation date", f"{ctx.valuation_date:%d %b %Y}"),
                ("Base currency", ctx.portfolio.base_currency),
                ("Positions", len(ctx.portfolio.positions)),
                ("Portfolio NAV (USD)", round(nav, 2)),
                (f"{cfg.holding_period}D VaR 99% historical (USD)", round(hist99.var, 2)),
                (f"{cfg.holding_period}D ES 99% historical (USD)", round(hist99.expected_shortfall, 2)),
                ("VaR 99% / NAV", round(hist99.var_pct, 6)),
                ("ES 99% / NAV", round(hist99.es_pct, 6)),
                (f"Portfolio volatility (annualised, {cfg.covariance_method})", round(portfolio_volatility, 6)),
                ("Worst stress scenario", worst.scenario.name if worst else "n/a"),
                ("Worst stress P&L (USD)", round(worst.total_pnl, 2) if worst else None),
                ("Lookback / decay factor", f"{cfg.lookback} days / {cfg.decay_factor:.2f}"),
            ],
            columns=["Metric", "Value"],
        )
        metrics = []
        for m in METHODS:
            for c, r in sorted(var_report.results[m].items()):
                metrics.append(
                    {
                        "Method": METHOD_LABELS[m], "Confidence": c, "Holding period": r.holding_period, "Lookback": r.lookback,
                        "Decay factor": r.decay_factor, "VaR (USD)": r.var, "VaR % NAV": r.var_pct, "ES (USD)": r.expected_shortfall,
                        "ES % NAV": r.es_pct, "ES / VaR": r.es_var_ratio, "1D sigma (USD)": r.volatility,
                    }
                )
        risk_metrics = pd.DataFrame(metrics)
        es = risk_metrics[risk_metrics["Confidence"] == 0.99][["Method", "VaR (USD)", "ES (USD)", "ES / VaR", "ES % NAV"]].reset_index(drop=True)
        var_cmp = var_report.comparison_table((0.95, 0.99), 0.99).reset_index(names="Measure")

        opt = sensitivities["options"]
        greeks = (
            opt[["position_id", "underlying", "option_type", "quantity", "strike", "expiry", "implied_vol", "market_value", "delta", "delta_usd", "gamma", "vega", "theta", "rho"]]
            if not opt.empty else pd.DataFrame()
        )
        bonds = sensitivities["bonds"]
        krd = sensitivities["krd"].reset_index(names="position_id") if not sensitivities["krd"].empty else pd.DataFrame()

        stress = pd.DataFrame([r.summary() | {"vol_method": r.vol_method, "bond_method": r.bond_method} for r in stress_results])
        if not stress.empty:
            stress["worst_asset_class"] = stress["worst_asset_class"].map(lambda c: ASSET_LABELS.get(str(c), c))
        pnl = summarise_components(attribution)
        pnl.insert(0, "date", attribution.date_to.date())
        pnl_hist = attribution_history.rename(columns=COMPONENT_LABELS).reset_index()

        bt = backtest.frame.reset_index(names="date")
        k = backtest.kupiec
        bt_summary = pd.DataFrame(
            [
                ("Model", METHOD_LABELS.get(backtest.method, backtest.method)),
                ("Confidence", backtest.confidence),
                ("Lookback", backtest.lookback),
                ("Period", f"{backtest.frame.index[0]:%d %b %Y} - {backtest.frame.index[-1]:%d %b %Y}"),
                ("Observations", k.n_observations),
                ("Exceptions", k.n_exceptions),
                ("Exception rate", round(k.exception_rate, 6)),
                ("Expected exception rate", round(k.expected_rate, 6)),
                ("Kupiec LR statistic", round(k.lr_statistic, 4)),
                ("Kupiec p-value", round(k.p_value, 4)),
                ("Kupiec result", k.message),
                *(
                    [
                        ("Christoffersen independence LR", round(backtest.christoffersen.lr_independence, 4)),
                        ("Independence p-value", round(backtest.christoffersen.p_value_independence, 4)),
                        ("Independence result", backtest.christoffersen.independence_message),
                        ("Conditional coverage LR", round(backtest.christoffersen.lr_conditional_coverage, 4)),
                        ("Conditional coverage p-value", round(backtest.christoffersen.p_value_conditional_coverage, 4)),
                    ]
                    if backtest.christoffersen is not None else []
                ),
                *(
                    [
                        ("Basel zone (99%, last 250 days)", f"{basel.zone.title()} ({basel.n_exceptions} exceptions)"),
                        ("Basel plus factor", "n/a" if basel.plus_factor is None else f"{basel.plus_factor:.2f}"),
                    ]
                    if basel is not None else []
                ),
            ],
            columns=["Metric", "Value"],
        )
        limits = pd.DataFrame(
            [{"Limit": c.name, "Value": c.value, "Limit level": c.limit, "Utilisation": c.utilisation, "Status": c.status} for c in limit_checks]
        )
        return RiskReportData(
            portfolio_id=ctx.portfolio.portfolio_id, portfolio_name=ctx.portfolio.name, valuation_date=ctx.valuation_date,
            base_currency=ctx.portfolio.base_currency, generated_at=datetime.now(), summary=summary,
            positions=ctx.valuation.drop(columns=["instrument_id"], errors="ignore"), risk_metrics=risk_metrics, var_comparison=var_cmp,
            expected_shortfall=es, contributions=contributions, contributions_by_class=contributions_by_class.reset_index(),
            greeks=greeks, bond_sensitivities=bonds, krd=krd, stress=stress, pnl_attribution=pnl, pnl_history=pnl_hist,
            backtest=bt, backtest_summary=bt_summary, limits=limits, warnings=list(concentration.warnings),
        )

    # ------------------------------------------------------------------- CSV
    @staticmethod
    def csv_exports(data: RiskReportData) -> dict[str, bytes]:
        names = {
            "Position risk": ("position_risk.csv", data.contributions),
            "VaR": ("var_comparison.csv", data.var_comparison),
            "ES": ("expected_shortfall.csv", data.expected_shortfall),
            "Risk metrics (all methods)": ("risk_metrics.csv", data.risk_metrics),
            "Stress results": ("stress_results.csv", data.stress),
            "Greeks": ("greeks.csv", data.greeks),
            "Bond sensitivities": ("bond_sensitivities.csv", data.bond_sensitivities),
            "KRD": ("key_rate_durations.csv", data.krd),
            "P&L attribution": ("pnl_attribution_history.csv", data.pnl_history),
            "Backtest results": ("var_backtest.csv", data.backtest),
        }
        return {label: (fname, df.to_csv(index=False).encode()) for label, (fname, df) in names.items()}

    # ----------------------------------------------------------------- Excel
    @staticmethod
    def excel(data: RiskReportData) -> bytes:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        buf = io.BytesIO()
        header_fill = PatternFill("solid", fgColor="0F1B30")
        header_font = Font(bold=True, color="FFFFFF")
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            for sheet, df in data.tables().items():
                frame = df if not df.empty else pd.DataFrame({"Note": ["No data"]})
                frame = frame.copy()
                for col in frame.columns:
                    if pd.api.types.is_datetime64_any_dtype(frame[col]):
                        frame[col] = frame[col].dt.date
                frame.to_excel(writer, sheet_name=sheet[:31], index=False, startrow=2)
                ws = writer.sheets[sheet[:31]]
                ws["A1"] = f"{sheet} - {data.portfolio_name} - valuation {data.valuation_date:%d %b %Y}"
                ws["A1"].font = Font(bold=True, size=12)
                for cell in ws[3]:
                    cell.fill, cell.font = header_fill, header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for i, col in enumerate(frame.columns, start=1):
                    width = max(len(str(col)), *(len(str(v)) for v in frame[col].head(200))) if len(frame) else len(str(col))
                    ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 48)
                    lname = str(col).lower()
                    fmt = None
                    if any(t in lname for t in ("usd", "market_value", "pnl", "p&l", "var", "exposure", "dv01", "nav", "loss")) and "pct" not in lname and "%" not in lname:
                        fmt = "#,##0.00;[Red](#,##0.00)"
                    if any(t in lname for t in ("pct", "% ", "weight", "rate", "ytm", "vol", "utilisation")) or lname.endswith("%"):
                        fmt = "0.00%"
                    if "marginal" in lname or "krd" in lname or "duration" in lname:
                        fmt = "0.0000"
                    if fmt:
                        for row in ws.iter_rows(min_row=4, min_col=i, max_col=i):
                            for cell in row:
                                if isinstance(cell.value, (int, float)):
                                    cell.number_format = fmt
                ws.freeze_panes = "A4"
            if data.warnings:
                pd.DataFrame({"Concentration warnings": data.warnings}).to_excel(writer, sheet_name="Warnings", index=False)
        logger.info("Excel risk report generated")
        return buf.getvalue()

    # ------------------------------------------------------------------- PDF
    @staticmethod
    def pdf(data: RiskReportData) -> bytes:
        from reportlab.graphics.charts.barcharts import HorizontalBarChart
        from reportlab.graphics.shapes import Drawing
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        navy = colors.HexColor("#0F1B30")
        grid = colors.HexColor("#D5D9E0")
        red, green = colors.HexColor("#B42828"), colors.HexColor("#0A7A0A")
        styles = getSampleStyleSheet()
        h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, textColor=navy, spaceAfter=2)
        h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, textColor=navy, spaceBefore=8, spaceAfter=4)
        body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=8.5, leading=11)
        small = ParagraphStyle("s", parent=body, fontSize=7.5, textColor=colors.HexColor("#555555"))

        def money(v) -> str:
            return "" if v is None or pd.isna(v) else f"{v:,.0f}"

        def pct(v) -> str:
            return "" if v is None or pd.isna(v) else f"{v * 100:.2f}%"

        def table(rows: list[list], widths: list[float] | None = None, signed_cols: tuple[int, ...] = ()) -> Table:
            t = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
            st = [
                ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, grid),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
            for r_i, row in enumerate(rows[1:], start=1):
                for c_i in signed_cols:
                    txt = str(row[c_i])
                    if txt.startswith("-") or txt.startswith("("):
                        st.append(("TEXTCOLOR", (c_i, r_i), (c_i, r_i), red))
            t.setStyle(TableStyle(st))
            return t

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                                title=f"Risk summary {data.portfolio_name}", author="Multi-Asset Risk Engine")
        s: list = [
            Paragraph("Multi-Asset Risk Engine - Executive Risk Summary", h1),
            Paragraph(f"{data.portfolio_name} ({data.portfolio_id}) · valuation {data.valuation_date:%d %b %Y} · base {data.base_currency} · "
                      f"generated {data.generated_at:%d %b %Y %H:%M}", small),
            Spacer(1, 6),
        ]
        s.append(Paragraph("1. Portfolio Summary", h2))
        s.append(table([["Metric", "Value"]] + [[m, (f"{v:,.2f}" if isinstance(v, float) and abs(v) > 1 else (pct(v) if isinstance(v, float) else str(v)))]
                                               for m, v in data.summary.itertuples(index=False)], [95 * mm, 80 * mm]))
        s.append(Paragraph("2. Risk Metrics", h2))
        rm = data.risk_metrics[data.risk_metrics["Confidence"].isin([0.95, 0.99])]
        s.append(table([["Method", "Conf.", "VaR (USD)", "VaR % NAV", "ES (USD)", "ES % NAV", "ES/VaR"]] +
                       [[r.Method, f"{r.Confidence:.0%}", money(r._6), pct(r._7), money(r._8), pct(r._9), f"{r._10:.2f}"] for r in rm.itertuples()]))
        s.append(Paragraph("3. VaR Comparison", h2))
        vc = data.var_comparison
        s.append(table([list(vc.columns)] + [[r[0]] + [money(v) for v in r[1:]] for r in vc.itertuples(index=False)]))
        s.append(Paragraph("4. Expected Shortfall", h2))
        s.append(Paragraph("VaR is the loss threshold exceeded in about 1% of scenarios at 99%; Expected Shortfall is the average "
                           "loss conditional on being in that worst 1% tail.", body))
        es = data.expected_shortfall
        s.append(table([["Method", "VaR 99%", "ES 99%", "ES/VaR", "ES % NAV"]] +
                       [[r.Method, money(r._2), money(r._3), f"{r._4:.2f}", pct(r._5)] for r in es.itertuples()]))
        s.append(Paragraph("5. Risk Contributions (parametric Euler allocation)", h2))
        rc = data.contributions.sort_values("component_var", ascending=False)
        s.append(table([["Position", "Ticker", "Class", "Exposure", "Marginal VaR", "Component VaR", "% VaR"]] +
                       [[r.position_id, r.ticker, r.asset_class, money(r.exposure), f"{r.marginal_var:.4f}", money(r.component_var), pct(r.pct_contribution)]
                        for r in rc.itertuples()], signed_cols=(5, 6)))
        s.append(PageBreak())
        s.append(Paragraph("6. Sensitivities", h2))
        if not data.greeks.empty:
            g = data.greeks
            s.append(table([["Position", "Option", "IV", "MV", "Delta $", "Gamma", "Vega/pt", "Theta/day", "Rho/1%"]] +
                           [[r.position_id, f"{r.underlying} {r.option_type[0]}{r.strike:g}", pct(r.implied_vol), money(r.market_value),
                             money(r.delta_usd), f"{r.gamma:,.1f}", money(r.vega), money(r.theta), money(r.rho)] for r in g.itertuples()],
                           signed_cols=(3, 4, 6, 7, 8)))
            s.append(Spacer(1, 4))
        if not data.bond_sensitivities.empty:
            b = data.bond_sensitivities
            s.append(table([["Position", "Bond", "MV", "Yield", "Duration", "Mod. Dur.", "Convexity", "DV01"]] +
                           [[r.position_id, r.ticker, money(r.market_value), pct(r.ytm), f"{r.macaulay_duration:.2f}", f"{r.modified_duration:.2f}",
                             f"{r.convexity:.1f}", money(r.dv01)] for r in b.itertuples()]))
        s.append(Paragraph("7. Stress Tests", h2))
        st_df = data.stress
        s.append(table([["Scenario", "P&L (USD)", "P&L %", "Stressed NAV", "Worst position", "Worst class"]] +
                       [[r.scenario, money(r.total_pnl), pct(r.pnl_pct), money(r.stressed_nav), r.worst_position, r.worst_asset_class]
                        for r in st_df.itertuples()], signed_cols=(1, 2)))
        if not st_df.empty:
            d = Drawing(170 * mm, 60 * mm)
            ch = HorizontalBarChart()
            ch.x, ch.y, ch.width, ch.height = 45 * mm, 5 * mm, 115 * mm, 52 * mm
            ch.data = [[float(v) / 1e6 for v in st_df["total_pnl"]]]
            ch.categoryAxis.categoryNames = list(st_df["scenario"])
            ch.categoryAxis.labels.fontSize = 6.5
            ch.valueAxis.labels.fontSize = 6.5
            ch.categoryAxis.labels.fontName = ch.valueAxis.labels.fontName = "Helvetica"
            ch.valueAxis.labelTextFormat = "%.1fM"
            ch.bars[0].fillColor = colors.HexColor("#3987E5")
            for i, v in enumerate(st_df["total_pnl"]):
                ch.bars[(0, i)].fillColor = red if v < 0 else green
            d.add(ch)
            s.append(Spacer(1, 4))
            s.append(d)
        s.append(Paragraph("8. P&amp;L Attribution", h2))
        pa = data.pnl_attribution
        s.append(table([["Component", "P&L (USD)"]] + [[r.component, f"{r.pnl:,.2f}"] for r in pa.itertuples()], [80 * mm, 40 * mm], signed_cols=(1,)))
        s.append(Paragraph("9. VaR Backtesting", h2))
        cell = ParagraphStyle("cell", parent=body, fontSize=7.5, leading=9, alignment=2)

        def bt_value(m: str, v):
            if "rate" in m.lower() or m == "Confidence":
                return pct(v)
            text = str(v)
            return Paragraph(escape(text), cell) if len(text) > 70 else text

        s.append(table([["Metric", "Value"]] + [[m, bt_value(m, v)] for m, v in data.backtest_summary.itertuples(index=False)],
                       [60 * mm, 115 * mm]))
        s.append(Paragraph("10. Risk Limit Breaches", h2))
        lim = data.limits
        lt = table([["Limit", "Value", "Limit", "Utilisation", "Status"]] +
                   [[r.Limit, pct(r.Value), pct(r._3), pct(r.Utilisation), r.Status] for r in lim.itertuples()])
        for i, st_ in enumerate(lim["Status"], start=1):
            lt.setStyle(TableStyle([("TEXTCOLOR", (4, i), (4, i), red if st_ == "Breach" else green)]))
        s.append(lt)
        breaches = lim[lim["Status"] == "Breach"]
        s.append(Paragraph("No limit breaches." if breaches.empty else f"{len(breaches)} limit breach(es): " + ", ".join(breaches["Limit"]), body))
        if data.warnings:
            s.append(Paragraph("Concentration warnings", h2))
            for w in data.warnings:
                s.append(Paragraph(f"• {escape(w)}", body))
        s.append(Spacer(1, 8))
        s.append(Paragraph("All figures are computed from the stored portfolio and market data by the risk engine. Synthetic sample data. "
                           "VaR/ES: 1-day hypothetical P&amp;L of the static portfolio; stress tests are deterministic instantaneous shocks.", small))
        doc.build(s)
        logger.info("PDF risk summary generated")
        return buf.getvalue()
